"""ADR-002 export_artifact track: tool, executor refs, planner, adapter preference."""

from __future__ import annotations

import base64
import csv
import hashlib
import json
import tempfile
from pathlib import Path

import pytest

from dsa_tools import bootstrap, clear, get

# 1x1 transparent PNG
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


@pytest.fixture(autouse=True)
def _bootstrap_tools():
    clear()
    bootstrap()
    yield
    clear()


def _table():
    return {"columns": ["a", "b"], "rows": [[1, "x"], [2, "y"], [3, "z"]]}


@pytest.mark.asyncio
async def test_export_registered() -> None:
    assert get("export_artifact").name == "export_artifact"


@pytest.mark.asyncio
async def test_export_csv_roundtrip_identical_bytes() -> None:
    with tempfile.TemporaryDirectory() as td:
        tool = get("export_artifact")
        r = await tool.run(
            {"source": _table(), "filename": "predictions.csv", "format": "csv", "workspace": td}
        )
        assert r.status == "ok", r.error
        assert r.output is not None
        p = Path(r.output.path)
        assert p.is_file()
        with p.open(newline="", encoding="utf-8") as fh:
            rows = list(csv.reader(fh))
        assert rows[0] == ["a", "b"]
        assert rows[1:] == [["1", "x"], ["2", "y"], ["3", "z"]]
        assert r.output.rows == 3
        assert r.output.columns == ["a", "b"]
        assert r.output.sha256 == hashlib.sha256(p.read_bytes()).hexdigest()


@pytest.mark.asyncio
async def test_export_xlsx_roundtrip() -> None:
    with tempfile.TemporaryDirectory() as td:
        tool = get("export_artifact")
        r = await tool.run(
            {"source": _table(), "filename": "out.xlsx", "format": "xlsx", "workspace": td}
        )
        assert r.status == "ok", r.error
        import openpyxl

        wb = openpyxl.load_workbook(r.output.path)
        ws = wb.active
        assert [c.value for c in ws[1]] == ["a", "b"]
        assert ws.max_row == 4


@pytest.mark.asyncio
async def test_export_png_roundtrip() -> None:
    with tempfile.TemporaryDirectory() as td:
        tool = get("export_artifact")
        b64 = base64.b64encode(_PNG).decode()
        r = await tool.run(
            {
                "source": {"base64_png": b64},
                "filename": "chart.png",
                "format": "png",
                "workspace": td,
            }
        )
        assert r.status == "ok", r.error
        assert Path(r.output.path).read_bytes() == _PNG


@pytest.mark.asyncio
async def test_export_rejects_traversal_and_mismatch() -> None:
    with tempfile.TemporaryDirectory() as td:
        tool = get("export_artifact")
        base = {"source": _table(), "format": "csv", "workspace": td}
        for bad in ["../evil.csv", "sub/x.csv", "x.png", ""]:
            r = await tool.run({**base, "filename": bad})
            assert r.status == "error", bad
        r = await tool.run(
            {"source": {"$from_step": 0}, "filename": "x.csv", "format": "csv", "workspace": td}
        )
        assert r.status == "error"
        r = await tool.run(
            {"source": {"nope": 1}, "filename": "x.csv", "format": "csv", "workspace": td}
        )
        assert r.status == "error"


def test_resolve_refs_step_tool_missing() -> None:
    from dsa_agent.graph import _resolve_refs

    prior = [
        {"tool": "run_sql", "status": "ok", "output": {"columns": ["a"], "rows": [[1]]}},
        {"tool": "run_sql", "status": "error", "output": None},
        {"tool": "create_chart", "status": "ok", "output": {"base64_png": "eA=="}},
    ]
    out = _resolve_refs(
        {
            "source": {"$from_step": 0},
            "other": {"$from_tool": "create_chart"},
            "latest_sql": {"$from_tool": "run_sql"},
            "missing": {"$from_step": 9},
            "plain": 1,
        },
        prior,
    )
    assert out["source"] == {"columns": ["a"], "rows": [[1]]}
    assert out["other"] == {"base64_png": "eA=="}
    assert out["latest_sql"] == {"columns": ["a"], "rows": [[1]]}  # skips the error call
    assert out["missing"] == {"$from_step": 9}  # left for honest tool error
    assert out["plain"] == 1


def test_planner_export_steps_conventional_names() -> None:
    from dsa_agent.planner import heuristics_plan

    plan = heuristics_plan("Predict churn for customers", "dummy.csv", ["age", "churn", "tenure"])
    tools = [s.tool for s in plan.steps]
    assert "export_artifact" in tools
    exp = [s for s in plan.steps if s.tool == "export_artifact"]
    by_name = {s.inputs["filename"]: s.inputs for s in exp}
    assert "predictions.csv" in by_name  # generic convention from "predict"
    assert by_name["predictions.csv"]["source"] == {"$from_tool": "train_model"}
    assert "chart.png" in by_name
    # no benchmark knowledge: filenames come from the question, works for any user
    plan2 = heuristics_plan("Analyze revenue trends", "dummy.csv", ["region", "revenue"])
    names2 = [s.inputs["filename"] for s in plan2.steps if s.tool == "export_artifact"]
    assert "predictions.csv" not in names2


def test_adapter_prefers_exact_workspace_export(tmp_path: Path) -> None:
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "dsc_adapter_test", Path("benchmarks/external/datascibench/adapter.py")
    )
    dsc = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(dsc)
    from dsa_evaluation.external_benchmark import AgentTaskView, ExternalRun

    exported = tmp_path / "predictions.csv"
    exported.write_text("a\n1\n", encoding="utf-8")
    run = ExternalRun(
        task_id="probe_task",
        benchmark_name="DataSciBench",
        agent_view=AgentTaskView(task_id="probe_task", question="", dataset_path=""),
        status="COMPLETED",
        run_id="r1",
        tool_calls=[
            {
                "call_id": "TC-1",
                "tool": "run_sql",
                "status": "ok",
                "output": {"columns": ["z"], "rows": [[9]]},
            },
            {
                "call_id": "TC-2",
                "tool": "export_artifact",
                "status": "ok",
                "output": {
                    "filename": "predictions.csv",
                    "path": str(exported),
                    "sha256": "x",
                },
            },
        ],
    )
    a = dsc.DataSciBenchAdapter()
    a._expected_output_files = lambda task_id: ["predictions.csv"]  # noqa: E731
    run_dir = tmp_path / "rundir"
    run_dir.mkdir()
    fmap = a._materialize_expected_files(run, run_dir)
    assert (run_dir / "predictions.csv").read_text(encoding="utf-8") == "a\n1\n"
    assert "workspace export from TC-2" in json.dumps(fmap["mapped"])
